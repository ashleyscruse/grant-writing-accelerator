#!/usr/bin/env python3
"""Context Matters -- Thursday AI workshop deck (simplified). Ashley Scruse brand."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

COTTON = RGBColor(0xE5, 0xE6, 0xE0)
BERRY  = RGBColor(0x9E, 0x3E, 0x58)
BLOOD  = RGBColor(0xB3, 0x00, 0x22)
NOIR   = RGBColor(0x18, 0x17, 0x1F)
CDIM   = RGBColor(0xB9, 0xBA, 0xB4)
GRAY   = RGBColor(0x6B, 0x6B, 0x6B)
WHITE  = RGBColor(0xFA, 0xFA, 0xF8)
HEAD, SCRIPT, BODY = "Space Grotesk", "Caveat", "Outfit"
W, H = 13.333, 7.5

prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)

def _spc(run, pts): run.font._rPr.set('spc', str(int(pts*100)))

def slide(bg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background(); r.shadow.inherit = False
    return s

def rect(s, l, t, w, h, fill, line=None, lw=1.0):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb = line; r.line.width = Pt(lw)
    r.shadow.inherit = False
    return r

def text(s, txt, l, t, w, h, font=BODY, size=24, bold=False, color=NOIR,
         align=PP_ALIGN.LEFT, sp=1.1, track=None, italic=False):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(txt.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = sp
        run = p.add_run(); run.text = line
        run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
        run.font.italic = italic; run.font.color.rgb = color
        if track is not None: _spc(run, track)
    return tb

def kicker(s, txt, l=0.9, t=0.72, color=BERRY):
    text(s, txt, l, t, 10, 0.7, font=SCRIPT, size=32, bold=True, color=color)

def headline(s, txt, l=0.9, t=1.35, size=44, color=NOIR):
    text(s, txt, l, t, 11.6, 1.7, font=HEAD, size=size, bold=True, color=color, sp=1.02, track=-1.5)

def stamp(s, txt, l=0.9, t=0.32, color=GRAY):
    text(s, txt, l, t, 10, 0.35, font=HEAD, size=12, bold=True, color=color, track=2.5)

def footer(s, color=CDIM):
    # Website intentionally omitted for the workshop; add back after.
    return

def bullets(s, pairs, l=0.9, t=3.0, size=24, lead=BLOOD, rest=NOIR):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(11.6), Inches(4.0))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (a, b) in enumerate(pairs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12; p.space_after = Pt(15)
        r1 = p.add_run(); r1.text = a
        r1.font.name = HEAD; r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = lead; _spc(r1, -0.5)
        if b:
            r2 = p.add_run(); r2.text = "  " + b
            r2.font.name = BODY; r2.font.size = Pt(size); r2.font.color.rgb = rest

def cards(s, items, top=2.75):
    n = len(items); gap = 0.5; cw = (11.533 - gap*(n-1))/n
    for i, (num, title, desc) in enumerate(items):
        x = 0.9 + i*(cw+gap)
        rect(s, x, top, cw, 3.5, WHITE)
        text(s, num, x+0.28, top+0.25, 1.5, 0.5, font=HEAD, size=22, bold=True, color=BLOOD)
        rect(s, x+0.28, top+0.95, min(cw-0.56, 2.6), 0.42, BERRY)
        text(s, title, x+0.4, top+0.99, cw-0.6, 0.4, font=HEAD, size=13, bold=True, color=WHITE, track=1.0)
        text(s, desc, x+0.28, top+1.6, cw-0.5, 1.7, font=BODY, size=14.5, color=GRAY, sp=1.12)

def twocol(s, ll, lb, rl, rb, top=2.9):
    rect(s, 6.62, top, 0.035, 3.1, BERRY)
    text(s, ll, 0.9, top, 5.3, 0.6, font=HEAD, size=24, bold=True, color=BLOOD, track=-0.5)
    text(s, lb, 0.9, top+0.85, 5.3, 2.6, font=BODY, size=19, color=NOIR, sp=1.18)
    text(s, rl, 6.95, top, 5.3, 0.6, font=HEAD, size=24, bold=True, color=BLOOD, track=-0.5)
    text(s, rb, 6.95, top+0.85, 5.3, 2.6, font=BODY, size=19, color=NOIR, sp=1.18)

# ---- 1 COVER ----
s = slide(NOIR); stamp(s, "GRANT WRITING ACCELERATOR   |   AI WORKSHOP", color=CDIM)
kicker(s, "the one thing that matters", color=BERRY)
text(s, "CONTEXT MATTERS.", 0.9, 2.4, 12, 2.0, font=HEAD, size=76, bold=True, color=COTTON, track=-2.5)
rect(s, 0.95, 4.55, 2.4, 0.10, BLOOD)
text(s, "It is not the AI. It is what the AI knows about you.", 0.9, 4.95, 11, 0.8, font=BODY, size=22, color=CDIM)
footer(s)

# ---- 2 THE PROBLEM ----
s = slide(COTTON); stamp(s, "THE PROBLEM"); kicker(s, "here's the trap")
headline(s, "MOST SETUPS DON'T\nKEEP UP WITH YOU.", t=1.25)
text(s, "The web can hold some context: projects, uploaded PDFs. But it is a snapshot you keep up by hand, and it goes stale. A file-aware setup reads your real, current files every single time.",
     0.9, 3.7, 10.6, 2.2, font=BODY, size=22, color=NOIR, sp=1.25)
footer(s)

# ---- 3 THE WORDS ----
s = slide(COTTON); stamp(s, "SAME PAGE"); kicker(s, "learn three words")
headline(s, "THE WORDS.")
cards(s, [
    ("01", "CONTEXT", "What the AI can see right now: your prompt, your files, the chat so far."),
    ("02", "MEMORY", "It forgets between sessions. Your files are what give it memory."),
    ("03", "PROMPT", "How you ask. Clear, specific requests get better answers."),
])
footer(s)

# ---- 4 WHITEBOARD vs NOTEBOOK ----
s = slide(COTTON); stamp(s, "HOW MEMORY WORKS"); kicker(s, "where things live")
headline(s, "WHITEBOARD vs NOTEBOOK.", size=40)
twocol(s, "THE WHITEBOARD",
       "The session you are in. Temporary. Wiped clean when you close the window.",
       "THE NOTEBOOK",
       "Your files and CLAUDE.md. Permanent on disk. Reloaded onto the whiteboard every session.")
footer(s)

# ---- 5 WHAT FILE-AWARE AI MEANS ----
s = slide(COTTON); stamp(s, "THE WHOLE IDEA"); kicker(s, "this is the shift")
headline(s, "WHAT FILE-AWARE AI MEANS.", size=38)
twocol(s, "PASTE AND COPY",
       "You paste text into a chat and copy the answer back. It only ever sees what you paste.",
       "FILE-AWARE",
       "It opens your folder and reads and writes your real files, in place. It sees your whole project.")
footer(s)

# ---- 6 HOW IT SPEEDS YOU UP ----
s = slide(COTTON); stamp(s, "WHY IT MATTERS"); kicker(s, "the payoff")
headline(s, "HOW THAT SPEEDS YOU UP.", size=40)
bullets(s, [
    ("No re-explaining.", "It already knows you, from your files."),
    ("Drafts in place.", "It writes into your documents, not a chat you copy from."),
    ("Stays organized.", "Your work lives in folders, not scattered tabs."),
], t=3.05)
footer(s)

# ---- 7 MEET CLAUDE CODE ----
s = slide(COTTON); stamp(s, "THE FILE-AWARE TOOL"); kicker(s, "what you'll use")
headline(s, "MEET CLAUDE CODE.")
bullets(s, [
    ("It opens your folder.", "Reads and writes your files, in plain English."),
    ("You never touch code.", "You ask; it does the work."),
    ("It even writes the LaTeX.", "The formatting that keeps NSF pages compliant."),
], t=3.05)
text(s, "You'll use the desktop app. No coding editor required.",
     0.9, 6.35, 11.5, 0.6, font=BODY, size=16, italic=True, color=GRAY)
footer(s)

# ---- 8 CLAUDE PRO ----
s = slide(COTTON); stamp(s, "THE PLAN"); kicker(s, "what you'll sign up for")
headline(s, "ABOUT CLAUDE PRO.")
text(s, "$20 per month. Includes Claude Code.\n\nUsage resets on a rolling 5-hour window, plus a weekly cap.\n\nGood context stretches it: it reads your files instead of you re-typing.",
     0.9, 3.0, 7.4, 3.2, font=BODY, size=21, color=NOIR, sp=1.2)
text(s, "$20", 9.2, 2.9, 3.6, 1.6, font=HEAD, size=90, bold=True, color=BLOOD, align=PP_ALIGN.CENTER, track=-2)
text(s, "/ month", 9.2, 4.6, 3.6, 0.6, font=HEAD, size=22, bold=True, color=BERRY, align=PP_ALIGN.CENTER)
footer(s)

# ---- 9 WHAT'S IN THE WORKSPACE ----
s = slide(COTTON); stamp(s, "YOUR TURN SOON"); kicker(s, "what you're getting")
headline(s, "WHAT'S IN THE WORKSPACE.", size=40)
cards(s, [
    ("01", "PROFILE", "Who you are, built from your CV. Reused on every grant."),
    ("02", "GRANTS", "A folder per grant: solicitation, evaluation, budget, draft."),
    ("03", "SKILLS", "Sage, advisor, writer, reviewers. They draft and critique for you."),
], top=2.55)
text(s, "A tracker keeps your whole pipeline in view.",
     0.9, 6.35, 11.5, 0.6, font=BODY, size=16, italic=True, color=GRAY)
footer(s)

# ---- 10 WHAT STAYS HUMAN ----
s = slide(NOIR); stamp(s, "THE LINE I WON'T CROSS", color=CDIM); kicker(s, "some things aren't automated", color=BERRY)
text(s, "The budget stays human.\nYou and your team own it.", 0.9, 2.4, 11.6, 2.2,
     font=HEAD, size=40, bold=True, color=COTTON, sp=1.1, track=-1)
text(s, "That is why there is no budget skill. I never build a workflow that removes the person.",
     0.9, 5.0, 11.0, 1.0, font=BODY, size=20, color=CDIM, sp=1.2)
footer(s)

# ---- 11 CLOSE ----
s = slide(NOIR); stamp(s, "HANDS ON THE KEYBOARD", color=CDIM); kicker(s, "your turn.", color=BERRY)
text(s, "LET'S OPEN IT.", 0.9, 2.4, 12, 1.8, font=HEAD, size=72, bold=True, color=COTTON, track=-2.5)
rect(s, 0.95, 4.35, 2.4, 0.10, BLOOD)
text(s, "Download the desktop app and the zip. Unzip it. Open the folder in Claude Code. Say hi.",
     0.9, 4.75, 11.2, 1.0, font=BODY, size=22, color=CDIM, sp=1.2)
footer(s)

out = "context-matters-workshop-slides.pptx"
prs.save(out)
print("wrote", out, "with", len(prs.slides._sldIdLst), "slides")
