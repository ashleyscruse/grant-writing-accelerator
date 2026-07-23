#!/usr/bin/env python3
"""Context Matters -- LEAN demo intro (6 slides). Ashley Scruse brand."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

COTTON = RGBColor(0xE5,0xE6,0xE0); BERRY=RGBColor(0x9E,0x3E,0x58); BLOOD=RGBColor(0xB3,0x00,0x22)
NOIR=RGBColor(0x18,0x17,0x1F); CDIM=RGBColor(0xB9,0xBA,0xB4); GRAY=RGBColor(0x6B,0x6B,0x6B); WHITE=RGBColor(0xFA,0xFA,0xF8)
HEAD,SCRIPT,BODY="Space Grotesk","Caveat","Outfit"; W,H=13.333,7.5
prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)

def _spc(r,p): r.font._rPr.set('spc',str(int(p*100)))
def slide(bg):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(W),Inches(H))
    r.fill.solid(); r.fill.fore_color.rgb=bg; r.line.fill.background(); r.shadow.inherit=False; return s
def rect(s,l,t,w,h,fill,line=None,lw=1.0):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb=line; r.line.width=Pt(lw)
    r.shadow.inherit=False; return r
def text(s,txt,l,t,w,h,font=BODY,size=24,bold=False,color=NOIR,align=PP_ALIGN.LEFT,sp=1.1,track=None,italic=False):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    for i,line in enumerate(txt.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.line_spacing=sp
        run=p.add_run(); run.text=line; run.font.name=font; run.font.size=Pt(size); run.font.bold=bold
        run.font.italic=italic; run.font.color.rgb=color
        if track is not None: _spc(run,track)
    return tb
def kicker(s,txt,l=0.9,t=0.72,color=BERRY): text(s,txt,l,t,10,0.7,font=SCRIPT,size=32,bold=True,color=color)
def headline(s,txt,l=0.9,t=1.35,size=44,color=NOIR): text(s,txt,l,t,11.6,1.7,font=HEAD,size=size,bold=True,color=color,sp=1.02,track=-1.5)
def stamp(s,txt,l=0.9,t=0.32,color=GRAY): text(s,txt,l,t,10,0.35,font=HEAD,size=12,bold=True,color=color,track=2.5)
def bullets(s,pairs,l=0.9,t=3.0,size=24,lead=BLOOD,rest=NOIR):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(11.6),Inches(4.0)); tf=tb.text_frame; tf.word_wrap=True
    for i,(a,b) in enumerate(pairs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.line_spacing=1.12; p.space_after=Pt(15)
        r1=p.add_run(); r1.text=a; r1.font.name=HEAD; r1.font.size=Pt(size); r1.font.bold=True; r1.font.color.rgb=lead; _spc(r1,-0.5)
        if b:
            r2=p.add_run(); r2.text="  "+b; r2.font.name=BODY; r2.font.size=Pt(size); r2.font.color.rgb=rest
def cards(s,items,top=2.75):
    n=len(items); gap=0.5; cw=(11.533-gap*(n-1))/n
    for i,(num,title,desc) in enumerate(items):
        x=0.9+i*(cw+gap); rect(s,x,top,cw,3.5,WHITE)
        text(s,num,x+0.28,top+0.25,1.5,0.5,font=HEAD,size=22,bold=True,color=BLOOD)
        rect(s,x+0.28,top+0.95,min(cw-0.56,2.6),0.42,BERRY)
        text(s,title,x+0.4,top+0.99,cw-0.6,0.4,font=HEAD,size=13,bold=True,color=WHITE,track=1.0)
        text(s,desc,x+0.28,top+1.6,cw-0.5,1.7,font=BODY,size=14.5,color=GRAY,sp=1.12)
def twocol(s,ll,lb,rl,rb,top=2.9):
    rect(s,6.62,top,0.035,3.1,BERRY)
    text(s,ll,0.9,top,5.3,0.6,font=HEAD,size=24,bold=True,color=BLOOD,track=-0.5)
    text(s,lb,0.9,top+0.85,5.3,2.6,font=BODY,size=19,color=NOIR,sp=1.18)
    text(s,rl,6.95,top,5.3,0.6,font=HEAD,size=24,bold=True,color=BLOOD,track=-0.5)
    text(s,rb,6.95,top+0.85,5.3,2.6,font=BODY,size=19,color=NOIR,sp=1.18)

# 1 COVER
s=slide(NOIR); stamp(s,"GRANT WRITING ACCELERATOR   |   AI WORKSHOP",color=CDIM)
kicker(s,"before we build",color=BERRY)
text(s,"CONTEXT MATTERS.",0.9,2.4,12,2.0,font=HEAD,size=76,bold=True,color=COTTON,track=-2.5)
rect(s,0.95,4.55,2.4,0.10,BLOOD)
text(s,"It is not the AI. It is what it knows about you.",0.9,4.95,11,0.8,font=BODY,size=22,color=CDIM)

# 2 PAYOFF
s=slide(COTTON); stamp(s,"THE PAYOFF"); kicker(s,"why you're here")
headline(s,"WHAT YOU'LL WALK OUT WITH.",size=40)
bullets(s,[
 ("A workspace that knows you.","Your CV, your institution, your record."),
 ("Grant help in your voice.","Drafted to the review criteria, not generic mush."),
 ("Reusable on every grant.","Set it up once, use it for years."),
],t=3.05)

# 3 PROBLEM (concrete hook)
s=slide(COTTON); stamp(s,"SOUND FAMILIAR?"); kicker(s,"you've done this")
headline(s,"PASTE YOUR BIO,\nGET GENERIC MUSH.",t=1.25,size=44)
text(s,"You drop your bio into a chat box and get back something that sounds nothing like you. That is not the AI's fault. It just does not know who you are.",
     0.9,3.9,10.6,2.0,font=BODY,size=22,color=NOIR,sp=1.25)

# 4 THE FIX
s=slide(COTTON); stamp(s,"THE SHIFT"); kicker(s,"here's the fix")
headline(s,"LET IT READ YOUR FILES.",size=42)
twocol(s,"PASTE AND COPY","You paste text into a chat. It only ever sees what you paste.",
       "FILE-AWARE","It opens your folder, reads your real files, and knows you, without re-typing a thing.")

# 5 THE WORDS (grant-grounded)
s=slide(COTTON); stamp(s,"THE VOCABULARY"); kicker(s,"three quick words")
headline(s,"THE WORDS.")
cards(s,[
 ("01","CONTEXT","What it sees right now: your CV and the solicitation."),
 ("02","MEMORY","Your files, so you never re-explain yourself."),
 ("03","PROMPT","How you ask: “draft my Broader Impacts.”"),
])

# 6 LET'S OPEN IT
s=slide(NOIR); stamp(s,"HANDS ON THE KEYBOARD",color=CDIM); kicker(s,"your turn.",color=BERRY)
text(s,"LET'S OPEN IT.",0.9,2.4,12,1.8,font=HEAD,size=72,bold=True,color=COTTON,track=-2.5)
rect(s,0.95,4.35,2.4,0.10,BLOOD)
text(s,"Watch it read my CV and know me. Then you'll do it with yours.",
     0.9,4.75,11.2,1.0,font=BODY,size=22,color=CDIM,sp=1.2)

out="context-matters-intro.pptx"; prs.save(out); print("wrote",out,"with",len(prs.slides._sldIdLst),"slides")
